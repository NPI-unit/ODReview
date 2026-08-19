"""Cross-group overview workbook  ->  outputs/Uebersicht.xlsx

Aggregates every (timepoint, subject) group into one workbook:

  • Sheet "Daten"        – tidy, one row per (Zeitpunkt, Condition, Proband)
                           with counts + percentages per type, %verändert,
                           %verändert_2000, %Target.
  • Sheet "%verändert"   – matrix Proband × Zeitpunkt.
  • Sheet "%<Typ>"       – one matrix per cell type (aggregated %).
  • Sheet "RBC_nBilder"  – control matrix (Σ active cells + image count).

All numbers come from the same eval-capped aggregation (group_eval_summary)
that the per-group summary CSVs use, so the values match exactly.
"""

from __future__ import annotations

import re
from pathlib import Path

from .grouping import UNGROUPED_TIMEPOINT, UNGROUPED_SUBJECT, natural_key


def _split_tp(tp: str) -> tuple[str, str]:
    """'T1-4' -> ('T1','4');  'T0' -> ('T0','');  ungrouped -> (tp,'')."""
    if tp == UNGROUPED_TIMEPOINT:
        return (tp, "")
    if "-" in tp:
        a, b = tp.split("-", 1)
        return (a, b)
    return (tp, "")


def build_overview(folder: Path, profile, grouped) -> dict:
    """Compute all tables. Returns a dict consumed by write_xlsx() and the
    in-app preview. tidy_rows is empty when no group has cached masks."""
    from .evaluation import group_eval_summary

    type_keys_non_normal = [t.key for t in profile.cell_types if t.key != "normal"]
    type_keys_all = type_keys_non_normal + ["normal"]
    type_name = {t.key: (t.name or t.key) for t in profile.cell_types}
    type_name.setdefault("normal", "Normal")
    target = max(1, profile.target_cells)

    # Collect per-group aggregates in stable (timepoint, subject) order.
    records: list[tuple[str, str, dict]] = []
    if grouped is not None:
        for tp in grouped.timepoints():
            for sub in grouped.subjects_in(tp):
                summ = group_eval_summary(folder, grouped.tree[tp][sub], profile)
                if summ["n_images"] == 0:
                    continue  # no cached masks in this group
                records.append((tp, sub, summ))

    tidy_header = (
        ["Zeitpunkt", "Condition", "Proband", "nBilder", "RBC"]
        + [type_name.get(k, k) for k in type_keys_all]
        + ["%" + type_name.get(k, k) for k in type_keys_all]
        + ["%verändert", "%verändert_2000", "%Target", "Ziel"]
    )
    tidy_rows: list[list] = []
    for tp, sub, s in records:
        timepoint, cond = _split_tp(tp)
        rbc = s["n_active"]
        by = s["by_type"]
        row: list = [timepoint, cond, sub, s["n_images"], rbc]
        row += [by.get(k, 0) for k in type_keys_all]
        row += [round(100.0 * by.get(k, 0) / rbc, 2) if rbc > 0 else 0.0
                for k in type_keys_all]
        row += [round(s["pct_changed"], 2), round(s["pct_changed_2000"], 2),
                round(100.0 * rbc / target, 2), target]
        tidy_rows.append(row)

    # Matrix axes.
    timepoints = sorted({tp for tp, _, _ in records},
                        key=lambda k: (k == UNGROUPED_TIMEPOINT, natural_key(k)))
    subjects = sorted({sub for _, sub, _ in records},
                      key=lambda k: (k == UNGROUPED_SUBJECT, natural_key(k)))
    lookup = {(tp, sub): s for tp, sub, s in records}

    def matrix(valfn) -> list[list]:
        out = []
        for sub in subjects:
            r: list = [sub]
            for tp in timepoints:
                s = lookup.get((tp, sub))
                r.append(round(valfn(s), 2) if s is not None else None)
            out.append(r)
        return out

    matrices: dict[str, list[list]] = {"%verändert": matrix(lambda s: s["pct_changed"])}
    for k in type_keys_all:
        matrices["%" + type_name.get(k, k)] = matrix(
            lambda s, k=k: (100.0 * s["by_type"].get(k, 0) / s["n_active"])
            if s["n_active"] > 0 else 0.0)

    return {
        "tidy_header": tidy_header,
        "tidy_rows": tidy_rows,
        "timepoints": timepoints,
        "subjects": subjects,
        "matrices": matrices,
        "rbc_matrix": matrix(lambda s: s["n_active"]),
        "n_matrix": matrix(lambda s: s["n_images"]),
        "type_name": type_name,
        "type_keys_all": type_keys_all,
        "target": target,
    }


def _safe_title(title: str, used: set[str]) -> str:
    """Excel sheet titles: <=31 chars, no []:*?/\\, unique."""
    t = re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet"
    base, i = t, 1
    while t in used:
        suf = f"_{i}"
        t = base[:31 - len(suf)] + suf
        i += 1
    used.add(t)
    return t


def write_xlsx(data: dict, dest: Path) -> Path | None:
    """Write the workbook. Returns the path, or None if there is no data."""
    if not data["tidy_rows"]:
        return None
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    used: set[str] = set()
    wb = Workbook()

    # Sheet 1: Daten (tidy)
    ws = wb.active
    ws.title = _safe_title("Daten", used)
    ws.append(data["tidy_header"])
    for c in ws[1]:
        c.font = bold
    for row in data["tidy_rows"]:
        ws.append(row)
    ws.freeze_panes = "A2"

    def add_matrix(title: str, matrix: list[list]) -> None:
        s = wb.create_sheet(_safe_title(title, used))
        s.append(["Proband \\ Zeitpunkt"] + list(data["timepoints"]))
        for c in s[1]:
            c.font = bold
        for row in matrix:
            s.append(row)
        s.freeze_panes = "B2"

    add_matrix("%verändert", data["matrices"]["%verändert"])
    for k in data["type_keys_all"]:
        name = "%" + data["type_name"].get(k, k)
        add_matrix(name, data["matrices"][name])

    # Control sheet: RBC + nBilder stacked.
    sc = wb.create_sheet(_safe_title("RBC_nBilder", used))
    sc.append(["RBC (Σ aktiv)"]); sc["A1"].font = bold
    sc.append(["Proband \\ Zeitpunkt"] + list(data["timepoints"]))
    for c in sc[2]:
        c.font = bold
    for row in data["rbc_matrix"]:
        sc.append(row)
    sc.append([])
    hdr = sc.max_row + 1
    sc.append(["nBilder"]); sc[f"A{hdr}"].font = bold
    sc.append(["Proband \\ Zeitpunkt"] + list(data["timepoints"]))
    for c in sc[hdr + 1]:
        c.font = bold
    for row in data["n_matrix"]:
        sc.append(row)

    # Definitionen / legend — so a reader can see how every value is computed.
    from .db import METRIC_DEFINITIONS
    sdef = wb.create_sheet(_safe_title("Definitionen", used))
    sdef.append(["Größe", "Formel", "Beschreibung"])
    for c in sdef[1]:
        c.font = bold
    for _key, label, formula, desc in METRIC_DEFINITIONS:
        sdef.append([label, formula, desc])
    sdef.append([])
    sdef.append(["Ziel (target_cells) bei diesem Export", data.get("target")])
    sdef.cell(row=sdef.max_row, column=1).font = bold
    sdef.column_dimensions["A"].width = 28
    sdef.column_dimensions["B"].width = 44
    sdef.column_dimensions["C"].width = 72

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def write_overview_xlsx(folder: Path, profile, grouped, dest: Path) -> Path | None:
    """Build + write in one call. Returns path or None if nothing to write."""
    return write_xlsx(build_overview(folder, profile, grouped), dest)
